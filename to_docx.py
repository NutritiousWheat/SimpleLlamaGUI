import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def excluded(file_path):
    return (file_path.startswith("./build") or file_path.endswith(".clang-uml")
            or file_path.endswith(".gitmodules") or file_path.endswith("CMakeLists.txt.user")
            or file_path.startswith("./submodules") or file_path.startswith("./.venv")
            or file_path.startswith("./.github") or file_path.startswith("./.cache")
            or file_path.startswith("./docs") or file_path.startswith("./to_docx")
            or file_path.startswith("./.git") or file_path.endswith(".docx")
            or file_path.startswith("./.idea") or file_path.endswith(".docx#"))

def set_font(run, font_name, size_pt):
    run.font.name = font_name
    run.font.size = Pt(size_pt)


def add_files_to_docx(directory, output_docx):
    # Create a new Document
    doc = Document()

    # Walk through the directory recursively
    for root, dirs, files in os.walk(directory):
        for file_name in files:
            file_path = os.path.join(root, file_name)
            if excluded(file_path):
                continue

            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()

                # Add heading with filename
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(f'{file_path.replace("./", "")}')
                set_font(run, 'Times New Roman', 14)

                # Add content
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(content)
                set_font(run, 'Courier New', 14)

                # Optional page break after each file
                doc.add_page_break()

            except Exception as e:
                print(f"Could not read file: {file_path} | Error: {e}")

    # Save the final document
    doc.save(output_docx)
    print(f"All files have been saved to {output_docx}")

# Example usage
directory_to_scan = "."
output_file = "output_document.docx"
add_files_to_docx(directory_to_scan, output_file)